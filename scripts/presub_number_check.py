#!/usr/bin/env python3
"""Post PEFT-on-pSMP: verify key numbers appear in nar_wxd0801.docx."""
from __future__ import annotations

from pathlib import Path

from docx import Document

DOC = Path(__file__).resolve().parents[2] / "nar0801" / "nar_wxd0801.docx"

REQUIRED_SNIPPETS = [
    "0.078",
    "0.065",
    "p ≈ 0.011",  # spaced form in condensed Discussion
    "0.011",
    "0.036",
    "43/50",
    "0.012",
    "0.021",
    "0.029",  # Table 2 @50% FA pSMP
    "0.028",  # Temporal20_new pSMP FA (3dp; exact≈0.02849)
    "0.008",
    "0.017",
    "0.027",
    "0.011",
    "0.119",
    "0.006→0.080",
    "0.073",  # S18c / §3.8 cluster-disjoint Δ from unrounded FA
    "not significant (Table S9b)",
    "remains underpowered",
    "Table S18",
]
FORBIDDEN = [
    "essentially matching",
    "Temporal5 remains too small",
    "confirmatory support",
]
# "not confirmatory isolation" is allowed; bare positive claim is not
FORBIDDEN_POSITIVE = [
    ("confirmatory isolation", "not confirmatory isolation"),
]


def main() -> int:
    doc = Document(str(DOC))
    text = "\n".join(p.text for p in doc.paragraphs)

    missing = []
    for s in REQUIRED_SNIPPETS:
        if s not in text:
            missing.append(s)
    # Δ≈−0.013 optional after Discussion PEFT condensation (means + sign p retained)

    bad = [s for s in FORBIDDEN if s in text]
    for phrase, allow_if in FORBIDDEN_POSITIVE:
        stripped = text.replace(allow_if, "")
        if phrase in stripped:
            bad.append(phrase)

    abs_body = ""
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "Abstract":
            abs_body = doc.paragraphs[i + 1].text
            break
    abs_bad = [k for k in ("Temporal", "cluster-disjoint", "S18", "Boltz") if k in abs_body]

    print("DOC", DOC)
    print("missing", missing)
    print("forbidden_hits", bad)
    print("abstract_forbidden", abs_bad)
    return 1 if missing or bad or abs_bad else 0


if __name__ == "__main__":
    raise SystemExit(main())
