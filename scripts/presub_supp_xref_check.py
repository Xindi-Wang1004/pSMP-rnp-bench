#!/usr/bin/env python3
"""Cite vs caption vs local deposit for supplementary table IDs."""
from __future__ import annotations
import re
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
DOC = ROOT / "nar0801" / "nar_wxd0801.docx"
TAB = ROOT / "nar_paper_update" / "tables"

FORBIDDEN_ORPHANS = {"S2c", "S6d", "S6e"}  # known-retired / never shipped

def main() -> int:
    doc = Document(str(DOC))
    text = "\n".join(p.text for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                text += "\n" + c.text

    cited = set(re.findall(r"Tables?\s+S(\d+[a-d]?)", text))
    cited |= set(re.findall(r"Table S(\d+[a-d]?)", text))
    caps = set()
    for p in doc.paragraphs:
        m = re.match(r"Supplementary Table S(\d+[a-d]?)", p.text)
        if m:
            caps.add(m.group(1))
    deposit = set()
    for f in TAB.glob("TableS*"):
        m = re.match(r"TableS(\d+[a-d]?)", f.name)
        if m:
            deposit.add(m.group(1))

    bad = []
    for oid in FORBIDDEN_ORPHANS:
        key = oid[1:]  # 2c
        if key in cited:
            bad.append(f"retired id still cited: {oid}")
    missing_cap = sorted(cited - caps)
    # lettered companions may be deposit-only with stub captions; require caption OR deposit
    missing_both = sorted(cited - caps - deposit)
    print("cited", sorted(cited))
    print("missing_caption", missing_cap)
    print("missing_caption_and_deposit", missing_both)
    print("bad", bad)
    # soft: allow missing caption if deposit exists (stub should cover lettered)
    return 1 if bad or missing_both else 0

if __name__ == "__main__":
    raise SystemExit(main())
